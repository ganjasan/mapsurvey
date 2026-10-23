"""`rest-api-v1`: export parity and event, import jobs, layers."""
import gzip
from io import BytesIO
from unittest import mock

from django.contrib.gis.geos import Point
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext

from survey.models import Answer, Question, SurveyImportJob, SurveyMapLayer, SurveySession

from .helpers import LOCMEM, V1Client, make_pro_org, make_v1_survey, mint


def _xlsx_rows(content):
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(content), read_only=True)
    return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


@LOCMEM
class V1DataTest(TestCase):
    def setUp(self):
        self.org, self.owner = make_pro_org()
        self.survey, self.section = make_v1_survey(self.org, self.owner)
        q = Question.objects.create(survey_section=self.section, name='Where?', input_type='point',
                                    order_number=1, code='PT')
        session = SurveySession.objects.create(survey=self.survey)
        Answer.objects.create(survey_session=session, question=q, point=Point(13.4, 52.5))
        self.api = V1Client(self.client, mint(self.owner)[0])

    def test_export_matches_the_browser_download(self):
        """
        GIVEN a survey with one answered session
        WHEN it is exported as xlsx through the API and through download_data
        THEN both workbooks hold the same sheets and rows, and the API emits data_exported with surface api
        """
        with mock.patch('survey.product_events.emit') as emit:
            via_api = self.api.get(f'/surveys/{self.survey.uuid}/export', {'format': 'xlsx'})
        self.assertEqual(via_api.status_code, 200)
        api_bytes = b''.join(via_api.streaming_content)
        emit.assert_any_call('data_exported', self.owner.pk,
                             {'survey_id': str(self.survey.id), 'format': 'xlsx', 'surface': 'api'})
        self.client.force_login(self.owner)
        via_browser = self.client.get(f'/surveys/{self.survey.uuid}/download', {'format': 'xlsx'})
        self.assertEqual(_xlsx_rows(api_bytes), _xlsx_rows(b''.join(via_browser.streaming_content)))

    def test_unknown_format_is_400(self):
        """
        GIVEN an unknown export format
        WHEN it is requested
        THEN 400 validation on format
        """
        response = self.api.get(f'/surveys/{self.survey.uuid}/export', {'format': 'docx'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('format', response.json()['fields'])

    def test_import_job_is_the_callers_own(self):
        """
        GIVEN a ZIP posted through the API
        WHEN the job is read by its creator and by another member
        THEN 202 with the job; the creator sees it, the other member gets 404
        """
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .helpers import make_member
        with mock.patch('api.v1.data.run_survey_import.delay') as delay:
            response = self.client.post('/api/v1/surveys/import',
                                        {'file': SimpleUploadedFile('s.zip', b'PK\x05\x06' + b'\0' * 18)},
                                        HTTP_AUTHORIZATION=f'Bearer {self.api.token}')
        self.assertEqual(response.status_code, 202, response.content)
        job_id = response.json()['id']
        delay.assert_called_once_with(job_id)
        self.assertEqual(self.api.get(f'/import-jobs/{job_id}').status_code, 200)
        other = make_member(self.org, 'other_member')
        self.assertEqual(self.api.get(f'/import-jobs/{job_id}', token=mint(other)[0]).status_code, 404)
        self.assertEqual(SurveyImportJob.objects.get(pk=job_id).organization, self.org)

    def _layer(self):
        return SurveyMapLayer.objects.create(
            survey=self.survey, name='Parks', color='#00aa00', position=1,
            geojson='{"type":"FeatureCollection","features":[]}')

    def test_layer_list_never_selects_geometry(self):
        """
        GIVEN a reference layer
        WHEN the layers list is requested
        THEN no captured query selects geojson_gz
        """
        self._layer()
        with CaptureQueriesContext(connection) as ctx:
            response = self.api.get(f'/surveys/{self.survey.uuid}/layers')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()[0]['name'], 'Parks')
        layer_queries = [q['sql'] for q in ctx.captured_queries if 'survey_surveymaplayer' in q['sql']]
        self.assertTrue(layer_queries)
        self.assertFalse(any('geojson_gz' in sql for sql in layer_queries))

    def test_layer_geojson_gzip_passthrough(self):
        """
        GIVEN a reference layer
        WHEN its geojson is requested with Accept-Encoding gzip
        THEN the stored gzip bytes are sent with Content-Encoding gzip
        """
        layer = self._layer()
        response = self.api.get(f'/surveys/{self.survey.uuid}/layers/{layer.pk}/geojson',
                                HTTP_ACCEPT_ENCODING='gzip')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Encoding'], 'gzip')
        self.assertIn(b'FeatureCollection', gzip.decompress(response.content))
