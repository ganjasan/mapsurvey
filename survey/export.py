"""Responses export: one collector, several writers (spec responses-export-formats).

``collect()`` walks the database once and returns an ``ExportBundle``; every format is a
writer over that bundle, so the flat observations table, the Excel workbook and the GIS
files can never disagree with the GeoJSON the legacy archive holds -- they are the same
features, formatted differently.

The legacy GeoJSON+CSV ZIP (``format=zip``, the default) keeps its exact entry order and
content: creators hold old links and scripts, and the tests pin its values.
"""
import csv
import json
from survey.question_types import GEO_TYPES
import logging
import os
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from datetime import datetime, timezone as dt_timezone
from io import BytesIO
from xml.sax.saxutils import escape as xml_escape
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd

from . import other_option
from .models import Answer, FILE_INPUT_TYPES, Question, SurveySession

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Format catalogue
# ---------------------------------------------------------------------------

FORMATS = ('zip', 'xlsx', 'csv', 'gpkg', 'shp', 'kml')
OGR_FORMATS = frozenset({'gpkg', 'shp', 'kml'})
# Formats that come back as one file unless respondent uploads are requested.
SINGLE_FILE_FORMATS = frozenset({'xlsx', 'gpkg', 'kml'})

# GDAL is in the image for GeoDjango; the binary decides whether the GIS
# formats are offered at all. A dev machine without it keeps Excel and CSV.
OGR_AVAILABLE = shutil.which('ogr2ogr') is not None
OGR_TIMEOUT = 120

CONTENT_TYPES = {
    'zip': 'application/zip',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'gpkg': 'application/geopackage+sqlite3',
    'kml': 'application/vnd.google-earth.kml+xml',
}


class ExportFailed(Exception):
    """A writer could not produce its file; the message carries the tool's stderr."""


def _sanitize_filename(name):
    """Remove characters that are invalid in Windows filenames."""
    return re.sub(r'[<>:"/\\|?*]', '_', name)


# Every input type is classified for export. A type in none of these sets is a
# bug — it means INPUT_TYPE_CHOICES gained a member and nobody decided how it
# leaves the platform — so _answer_cell warns rather than dropping it silently,
# which is how datetime went missing from the download unnoticed.

# Carry respondent input; exported as a CSV column or a GeoJSON property.
EXPORT_VALUE_TYPES = frozenset({
    'text', 'text_line', 'number', 'range',
    'choice', 'rating', 'thumbs', 'multichoice', 'datetime', 'ranking',
    'photo', 'audio', 'document',
})

# Exported as GeoJSON layers in their own right, never as a cell.
EXPORT_GEOMETRY_TYPES = frozenset(GEO_TYPES)

# Presentational; they collect nothing, so there is nothing to export.
EXPORT_DISPLAY_ONLY_TYPES = frozenset({'image', 'html'})

# Returned by _answer_cell for questions that must not produce a column at all,
# which is distinct from a question that produces an empty one.
EXPORT_NO_COLUMN = object()

GEOMETRY_TYPE_NAMES = {'point': 'Point', 'line': 'LineString', 'polygon': 'Polygon',
                       'spraycan': 'MultiPoint'}
OGR_GEOMETRY_TYPES = {'point': 'wkbPoint', 'line': 'wkbLineString', 'polygon': 'wkbPolygon',
                      'spraycan': 'wkbMultiPoint'}
CRS84 = {"type": "name", "properties": {"name": "urn:ogc:def:crs:OGC:1.3:CRS84"}}

# Metadata every feature carries besides the sub-question answers.
FEATURE_META_KEYS = ('session', 'session_id', 'language', 'validation_status')
SHARED_MAP_KEYS = ('mark_key', 'votes_up', 'votes_down', 'comments')
# Spray clouds (spec spraycan-question): trailing column, present only when a
# spraycan question is in scope, so existing scripts keep their column indices.
SPRAY_KEYS = ('dot_count', 'area_m2')


def _upload_archive_path(question, answer):
    """Where a file answer lives inside the responses ZIP. The same string is
    the CSV/GeoJSON cell, so a row names the file sitting next to it."""
    return 'files/{sid}/{code}__{name}'.format(
        sid=answer.survey_session_id,
        code=question.code,
        name=_sanitize_filename(answer.upload.original_name),
    )


def _format_datetime_cell(raw):
    """Serialise a stored datetime answer as ISO 8601.

    Values that do not parse are passed through unchanged: a raw string the
    creator can still interpret beats a blank cell.
    """
    if not raw:
        return ""
    try:
        return datetime.fromisoformat(raw).isoformat()
    except (TypeError, ValueError):
        return raw


def _answer_cell(question, answers):
    """Format one question's answer for export.

    `answers` holds the rows belonging to this question and nothing else, which
    is what keeps a blank question from inheriting its neighbour's value: the
    result is computed per call rather than accumulated across a loop.

    Returns EXPORT_NO_COLUMN for questions that should not appear as a cell.
    """
    input_type = question.input_type

    if input_type in EXPORT_GEOMETRY_TYPES or input_type in EXPORT_DISPLAY_ONLY_TYPES:
        return EXPORT_NO_COLUMN

    if input_type not in EXPORT_VALUE_TYPES:
        logger.warning(
            "Export: question %s has unclassified input_type %r; exporting an "
            "empty column. Classify it in survey/export.py.",
            question.code, input_type,
        )
        return ""

    if not answers:
        # A flagged question keeps its adjacent write-in column on unanswered rows too.
        return other_option.export_cell(question, None, "")

    answer = answers[0]

    if input_type in ('photo', 'audio', 'document'):
        # Several files per question: the cell names every archive path. The
        # geo sub-answer path hands all rows in at once; the CSV loop passes
        # one at a time and concatenates in the caller — same separator.
        paths = [_upload_archive_path(question, a) for a in answers if a.upload_id]
        return '; '.join(paths)

    if input_type in ('text', 'text_line'):
        return answer.text if answer.text is not None else ""

    if input_type == 'datetime':
        return _format_datetime_cell(answer.text)

    if input_type in ('number', 'range'):
        if answer.numeric is not None:
            return answer.numeric
        if answer.selected_choices:
            return answer.selected_choices[0]
        return ""

    if input_type in ('choice', 'rating', 'thumbs'):
        names = answer.get_selected_choice_names()
        return other_option.export_cell(question, answer, names[0] if names else "")

    if input_type == 'multichoice':
        return other_option.export_cell(question, answer, "; ".join(answer.get_selected_choice_names()))

    if input_type == 'ranking':
        # One column per item, valued by its rank: a single "a > b > c" cell
        # reads well and analyses badly, and ranks are what the creator wants
        # to average.
        ranks = {}
        for position, code in enumerate(answer.selected_choices or [], start=1):
            ranks[f"{question.name}: {question.get_choice_name(code)}"] = position
        return ranks

    return ""


# ---------------------------------------------------------------------------
# The bundle
# ---------------------------------------------------------------------------

@dataclass
class GeoLayer:
    """One geo question's answers: the FeatureCollection the legacy ZIP writes,
    plus what the flat table needs and GeoJSON properties cannot carry."""
    question: Question
    name: str                      # sanitised question name, the file/layer base name
    collection: dict               # GeoJSON FeatureCollection
    geometries: list = field(default_factory=list)   # GEOS geometry per feature
    sessions: list = field(default_factory=list)     # SurveySession per feature
    sub_columns: list = field(default_factory=list)  # property keys from sub-questions, in order
    has_verdicts: bool = False


@dataclass
class ObjectSheet:
    """Answers ABOUT layer objects for one Objects-on-the-map question."""
    question: Question
    records: list


@dataclass
class ResultsLayer:
    """A layer's derived GeoJSON enriched with per-object aggregates."""
    layer_pk: int
    layer_name: str
    collection: dict


@dataclass
class VersionPart:
    """Everything exported for one version header, with its filename prefix."""
    survey: object
    prefix: str
    geo_layers: list = field(default_factory=list)
    object_sheets: list = field(default_factory=list)
    results_layers: list = field(default_factory=list)
    session_rows: list = field(default_factory=list)   # the legacy per-session CSV rows
    sessions: list = field(default_factory=list)       # SurveySession objects, same order
    file_answers: list = field(default_factory=list)


@dataclass
class ExportBundle:
    survey: object                 # the survey the URL named
    parts: list = field(default_factory=list)

    @property
    def has_verdicts(self):
        return any(layer.has_verdicts for part in self.parts for layer in part.geo_layers)

    @property
    def has_spray(self):
        return any(layer.question.input_type == 'spraycan' for part in self.parts for layer in part.geo_layers)


def collect(survey, version_surveys, excluded_session_ids=None):
    """Read every version in ``version_surveys`` (list of (header, prefix)) once."""
    excluded_session_ids = excluded_session_ids or set()
    bundle = ExportBundle(survey=survey)
    for header, prefix in version_surveys:
        bundle.parts.append(_collect_part(header, prefix, excluded_session_ids))
    return bundle


def _collect_part(survey, prefix, excluded_session_ids):
    from .object_stats import shared_map_verdicts

    part = VersionPart(survey=survey, prefix=prefix)

    for question in survey.geo_questions():
        layer_properties = {
            "survey": question.survey_section.survey_header.name,
            "survey_section": question.survey_section.name,
            "required": question.required,
        }
        layer = GeoLayer(question=question, name=_sanitize_filename(question.name), collection=None)
        features = []
        verdicts = shared_map_verdicts(survey, question)
        layer.has_verdicts = bool(verdicts)
        geo_type = question.input_type
        for geo_answer in question.answers():
            if geo_answer.survey_session_id in excluded_session_ids:
                continue

            geom = geo_answer.geometry
            if geom is None:
                continue
            if geo_type == "polygon":
                coordinates = [[[i[0], i[1]] for i in geom.coords[0]]]
            elif geo_type == "line":
                coordinates = [[i[0], i[1]] for i in geom.coords]
            elif geo_type == "spraycan":
                coordinates = [[i[0], i[1]] for i in geom.coords]
            else:
                coordinates = [geom.coords[0], geom.coords[1]]

            properties = {}
            for subquestion, rows in geo_answer.subAnswers().items():
                cell = _answer_cell(subquestion, rows)
                # Unlike the CSV, every sub-question keeps a property even when
                # it holds nothing: a feature collection whose attribute set
                # varies per feature is awkward to read in QGIS.
                if cell is EXPORT_NO_COLUMN:
                    properties[subquestion.name] = ""
                elif isinstance(cell, dict):
                    # Several columns for one answer: ranking, a write-in option.
                    properties.update(cell)
                else:
                    properties[subquestion.name] = cell
            for key in properties:
                if key not in layer.sub_columns:
                    layer.sub_columns.append(key)

            if geo_type == "spraycan":
                from survey import spray
                properties["dot_count"] = len(geom)
                properties["area_m2"] = int(round(spray.cloud_area_m2(geom)))

            session = geo_answer.survey_session
            properties["session"] = str(session)
            properties["session_id"] = geo_answer.survey_session_id
            properties["language"] = session.language or ''
            properties["validation_status"] = session.validation_status or ''
            # Shared map (spec shared-map-layer): the verdict other respondents
            # gave this mark rides on the author's own feature — ten residents
            # asking for the same corner are one feature with votes_up=9.
            if verdicts:
                properties.update(verdicts.get(geo_answer.pk) or {
                    "mark_key": "", "votes_up": 0, "votes_down": 0, "comments": 0})

            features.append({
                "type": "Feature",
                "properties": properties,
                "geometry": {"type": GEOMETRY_TYPE_NAMES[geo_type], "coordinates": coordinates},
            })
            layer.geometries.append(geom)
            layer.sessions.append(session)

        layer.collection = {
            "type": "FeatureCollection",
            "name": question.name,
            "crs": CRS84,
            "properties": layer_properties,
            "features": features,
        }
        part.geo_layers.append(layer)

    _collect_object_answers(part, survey, excluded_session_ids)

    for session in survey.sessions():
        if session.id in excluded_session_ids:
            continue
        properties = {}
        for answer in session.answers():
            if not answer.question:
                continue
            cell = _answer_cell(answer.question, [answer])
            # Geometry questions are exported as their own GeoJSON layers and
            # display-only questions collect nothing, so neither gets a column.
            if cell is EXPORT_NO_COLUMN:
                continue
            # A ranking answer is several columns, not one cell.
            if isinstance(cell, dict):
                properties.update(cell)
            elif (answer.question.input_type in FILE_INPUT_TYPES
                    and properties.get(answer.question.name)):
                # Several files on one question: one cell listing every path.
                properties[answer.question.name] += '; ' + cell
            else:
                properties[answer.question.name] = cell
        properties["session"] = str(session)
        properties["session_id"] = session.id
        properties["datetime"] = session.start_datetime
        properties["language"] = session.language or ''
        properties["validation_status"] = session.validation_status or ''
        part.session_rows.append(properties)
        part.sessions.append(session)

    # Respondent files ride along under files/<session_id>/, read through the
    # storage API so the same code serves filesystem and S3. Only attached
    # uploads of non-excluded sessions — orphans and trashed sessions stay out.
    part.file_answers = list(
        Answer.objects
        .filter(
            survey_session__survey=survey,
            question__input_type__in=FILE_INPUT_TYPES,
            upload__isnull=False,
        )
        .exclude(survey_session_id__in=excluded_session_ids)
        .select_related('question', 'upload')
    )
    return part


def _collect_object_answers(part, survey, excluded_session_ids):
    """Answers ABOUT layer objects (spec object-answers): one sheet per Objects-
    on-the-map question keyed by object, and the layer's derived GeoJSON
    enriched with per-object aggregates — never with free text."""
    from .object_stats import object_aggregates, flat_properties, sub_questions_of
    questions = Question.objects.filter(
        survey_section__survey_header=survey, input_type='layer_objects', layer__isnull=False,
    ).select_related('layer').order_by('survey_section_id', 'order_number')
    layers_done = set()
    for question in questions:
        subs = sub_questions_of(question)
        if not subs:
            # Display-only (spec layer-objects-question): nothing was asked, so
            # there is no sheet to write — the layer itself still gets its
            # results GeoJSON below when another question of this survey asks.
            continue
        rows = (Answer.objects
                .filter(question__in=subs, layer_object__isnull=False)
                .exclude(survey_session_id__in=excluded_session_ids)
                .select_related('layer_object', 'survey_session', 'question')
                .order_by('survey_session_id', 'layer_object__position', 'layer_object_id'))
        grouped = {}
        for a in rows:
            grouped.setdefault((a.survey_session_id, a.layer_object_id), []).append(a)
        records = []
        for (session_id, _obj_id), answers in grouped.items():
            session = answers[0].survey_session
            obj = answers[0].layer_object
            record = {
                'session_id': session_id,
                'object_key': obj.key,
                'object_title': obj.title,
                'object_category': obj.category,
            }
            by_q = {}
            for a in answers:
                by_q.setdefault(a.question, []).append(a)
            for sub_q in subs:
                cell = _answer_cell(sub_q, by_q.get(sub_q, []))
                if cell is EXPORT_NO_COLUMN:
                    continue
                if isinstance(cell, dict):
                    record.update(cell)
                else:
                    record[sub_q.name] = cell
            record['datetime'] = session.start_datetime
            record['language'] = session.language or ''
            record['validation_status'] = session.validation_status or ''
            if question.layer.source == 'question':
                record['status'] = obj.status   # the creator's moderation is part of their data
            records.append(record)
        part.object_sheets.append(ObjectSheet(question=question, records=records))

        layer = question.layer
        if layer.pk in layers_done:
            continue
        layers_done.add(layer.pk)
        aggregates = {}
        # Only THIS header's questions: the layer is shared with the draft copy
        # and the archived versions, whose sub-questions would otherwise add
        # empty columns to every feature.
        for q in layer.questions.filter(input_type='layer_objects', survey_section__survey_header=survey):
            for key, entry in object_aggregates(q, excluded_session_ids=excluded_session_ids).items():
                props = flat_properties(entry)
                aggregates.setdefault(key, {}).update(props)
        try:
            collection = json.loads(layer.geojson)
        except ValueError:
            continue
        for feature in collection.get('features') or []:
            key = (feature.get('properties') or {}).get('_key')
            feature.setdefault('properties', {}).update(aggregates.get(key, {'answers': 0}))
        collection['name'] = layer.name
        collection['crs'] = dict(CRS84)
        part.results_layers.append(ResultsLayer(layer_pk=layer.pk, layer_name=layer.name, collection=collection))


# ---------------------------------------------------------------------------
# Legacy writer: the GeoJSON + CSV archive, unchanged
# ---------------------------------------------------------------------------

def _geojson_bytes(collection):
    return json.dumps(collection, ensure_ascii=False).encode('utf8')


def write_legacy_zip(bundle, zip):
    """Today's archive, entry for entry. Objects sheets and results layers sit
    between the geo layers and the per-session CSV, as they always did."""
    for part in bundle.parts:
        prefix = part.prefix
        for layer in part.geo_layers:
            zip.writestr(prefix + layer.name + '.geojson', _geojson_bytes(layer.collection))
        results = {r.layer_pk: r for r in part.results_layers}
        written = set()
        for sheet in part.object_sheets:
            zip.writestr(prefix + 'objects_' + _sanitize_filename(sheet.question.code) + '.csv',
                         pd.DataFrame(sheet.records).to_csv())
            layer_pk = sheet.question.layer_id
            if layer_pk in results and layer_pk not in written:
                written.add(layer_pk)
                zip.writestr(prefix + 'layers/' + _sanitize_filename(results[layer_pk].layer_name) + '.results.geojson',
                             _geojson_bytes(results[layer_pk].collection))
        zip.writestr(prefix + _sanitize_filename(part.survey.name) + '.csv',
                     pd.DataFrame(part.session_rows).to_csv())
        for answer in part.file_answers:
            try:
                with answer.upload.file.open('rb') as stored:
                    zip.writestr(prefix + _upload_archive_path(answer.question, answer), stored.read())
            except Exception:
                logger.warning(
                    "Export: upload %s missing from storage; row keeps the path, file absent from ZIP.",
                    answer.upload_id,
                )


def build_legacy_zip(bundle):
    """The in-memory archive `download_data` has always returned."""
    in_memory = BytesIO()
    zip = ZipFile(in_memory, "a")
    write_legacy_zip(bundle, zip)
    # Windows bug fix
    for file in zip.filelist:
        file.create_system = 0
    zip.close()
    in_memory.seek(0)
    return in_memory


# ---------------------------------------------------------------------------
# Flat tables
# ---------------------------------------------------------------------------

OBSERVATION_FIXED_COLUMNS = (
    'session_id', 'session_start_utc', 'session_language', 'validation_status', 'version',
    'section', 'question', 'question_code', 'geometry_type', 'lat', 'lon', 'wkt',
)
SESSION_COLUMNS = (
    'session_id', 'session_start_utc', 'session_end_utc', 'language', 'validation_status',
    'opened_by_kind', 'version', 'tags', 'notes',
)


def _utc_naive(value):
    """Excel has no zone: every datetime leaves as naive UTC, and the column
    header says so. A naive value is taken as UTC (start_datetime's default)."""
    if value is None:
        return None
    if value.tzinfo is not None:
        value = value.astimezone(dt_timezone.utc)
    return value.replace(tzinfo=None, microsecond=0)


def observation_columns(bundle):
    """Fixed columns, then every sub-question column by NAME across all geo
    questions (a shared name is one column), then the shared-map columns."""
    columns = list(OBSERVATION_FIXED_COLUMNS)
    for part in bundle.parts:
        for layer in part.geo_layers:
            for key in layer.sub_columns:
                if key not in columns:
                    columns.append(key)
    if bundle.has_verdicts:
        columns.extend(SHARED_MAP_KEYS)
    if bundle.has_spray:
        columns.extend(SPRAY_KEYS)
    return columns


def observation_rows(bundle):
    """One row per placed feature. lat/lon are the point, or the centroid of a
    line/polygon/spray cloud; wkt carries the full geometry on every row."""
    for part in bundle.parts:
        for layer in part.geo_layers:
            question = layer.question
            for feature, geom, session in zip(layer.collection['features'], layer.geometries, layer.sessions):
                props = feature['properties']
                centre = geom if geom.geom_type == 'Point' else geom.centroid
                row = {
                    'session_id': session.id,
                    'session_start_utc': _utc_naive(session.start_datetime),
                    'session_language': session.language or '',
                    'validation_status': session.validation_status or '',
                    'version': part.survey.version_number,
                    'section': question.survey_section.name,
                    'question': question.name,
                    'question_code': question.code,
                    'geometry_type': feature['geometry']['type'],
                    'lat': centre.y,
                    'lon': centre.x,
                    'wkt': geom.wkt,
                }
                for key in layer.sub_columns:
                    row[key] = props.get(key, '')
                if layer.has_verdicts:
                    for key in SHARED_MAP_KEYS:
                        row[key] = props.get(key, '')
                if question.input_type == 'spraycan':
                    for key in SPRAY_KEYS:
                        row[key] = props.get(key, '')
                yield row


def responses_columns(bundle):
    """The legacy per-session CSV columns in order of first appearance, with
    the session start named for what it is."""
    columns = []
    for part in bundle.parts:
        for row in part.session_rows:
            for key in row:
                key = 'session_start_utc' if key == 'datetime' else key
                if key not in columns:
                    columns.append(key)
    return columns


def responses_rows(bundle):
    for part in bundle.parts:
        for row in part.session_rows:
            out = {}
            for key, value in row.items():
                if key == 'datetime':
                    out['session_start_utc'] = _utc_naive(value)
                else:
                    out[key] = value
            yield out


def sessions_rows(bundle):
    for part in bundle.parts:
        for session in part.sessions:
            yield {
                'session_id': session.id,
                'session_start_utc': _utc_naive(session.start_datetime),
                'session_end_utc': _utc_naive(session.end_datetime),
                'language': session.language or '',
                'validation_status': session.validation_status or '',
                'opened_by_kind': session.opened_by_kind,
                'version': part.survey.version_number,
                'tags': '; '.join(session.tags or []),
                'notes': session.notes or '',
            }


def object_sheets(bundle):
    """(title, columns, rows) per Objects-on-the-map question with answers."""
    for part in bundle.parts:
        for sheet in part.object_sheets:
            columns = []
            rows = []
            for record in sheet.records:
                out = {}
                for key, value in record.items():
                    if key == 'datetime':
                        key, value = 'session_start_utc', _utc_naive(value)
                    out[key] = value
                    if key not in columns:
                        columns.append(key)
                rows.append(out)
            title = part.prefix + 'objects_' + _sanitize_filename(sheet.question.code)
            yield title, columns, rows


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_SHEET_FORBIDDEN = re.compile(r'[\[\]:*?/\\]')


def _sheet_title(name, taken):
    title = _SHEET_FORBIDDEN.sub('_', name)[:31] or 'sheet'
    base, n = title, 2
    while title in taken:
        suffix = f'_{n}'
        title = base[:31 - len(suffix)] + suffix
        n += 1
    taken.add(title)
    return title


def _xlsx_sheet(wb, title, columns, rows, numeric_columns=()):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.freeze_panes = 'A2'
    for i, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(name)) + 2))
    header = []
    for name in columns:
        cell = WriteOnlyCell(ws, value=str(name))
        cell.font = Font(bold=True)
        header.append(cell)
    ws.append(header)
    for row in rows:
        cells = []
        for name in columns:
            value = row.get(name, '')
            if value == '' or value is None:
                cells.append(None)
            elif isinstance(value, datetime):
                cell = WriteOnlyCell(ws, value=_utc_naive(value))
                cell.number_format = 'yyyy-mm-dd hh:mm'
                cells.append(cell)
            elif name in numeric_columns and isinstance(value, float):
                cell = WriteOnlyCell(ws, value=value)
                cell.number_format = '0.000000'
                cells.append(cell)
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cells.append(value)
            elif isinstance(value, (list, dict)):
                cells.append(json.dumps(value, ensure_ascii=False))
            else:
                cells.append(str(value))
        ws.append(cells)


def write_xlsx(bundle, path):
    """observations, responses, sessions, objects_<code>… in write-only mode:
    one row in memory at a time, the bundle is what the legacy ZIP holds anyway."""
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    taken = set()
    _xlsx_sheet(wb, _sheet_title('observations', taken), observation_columns(bundle),
                observation_rows(bundle), numeric_columns=('lat', 'lon'))
    _xlsx_sheet(wb, _sheet_title('responses', taken), responses_columns(bundle), responses_rows(bundle))
    _xlsx_sheet(wb, _sheet_title('sessions', taken), list(SESSION_COLUMNS), sessions_rows(bundle))
    for title, columns, rows in object_sheets(bundle):
        _xlsx_sheet(wb, _sheet_title(title, taken), columns, rows)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# CSV (flat tables, Excel-friendly)
# ---------------------------------------------------------------------------

def _csv_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return _utc_naive(value).isoformat()
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_csv_entry(zip, arcname, columns, rows):
    """UTF-8 with BOM so Excel on Windows reads it as UTF-8; `.` decimal, ISO datetimes."""
    with zip.open(arcname, 'w') as raw:
        raw.write('﻿'.encode('utf-8'))
        text = _Utf8Writer(raw)
        writer = csv.DictWriter(text, fieldnames=columns, extrasaction='ignore', lineterminator='\n')
        writer.writeheader()
        for row in rows:
            writer.writerow({k: _csv_value(row.get(k, '')) for k in columns})


class _Utf8Writer:
    """csv.writer wants text; ZipFile.open('w') wants bytes."""
    def __init__(self, raw):
        self.raw = raw

    def write(self, s):
        self.raw.write(s.encode('utf-8'))


def write_csv_zip(bundle, zip):
    _write_csv_entry(zip, 'observations.csv', observation_columns(bundle), observation_rows(bundle))
    _write_csv_entry(zip, 'responses.csv', responses_columns(bundle), responses_rows(bundle))
    _write_csv_entry(zip, 'sessions.csv', list(SESSION_COLUMNS), sessions_rows(bundle))
    for title, columns, rows in object_sheets(bundle):
        _write_csv_entry(zip, title + '.csv', columns, rows)


# ---------------------------------------------------------------------------
# GIS formats through ogr2ogr
# ---------------------------------------------------------------------------

def _unique_layer_names(bundle):
    """(name, collection, ogr geometry type) for every geo layer and results
    layer, names deduplicated with a numeric suffix. Two questions may share a
    name across sections; two versions always do."""
    taken = set()
    out = []

    def claim(base):
        name, n = base, 2
        while name in taken:
            name = f'{base}_{n}'
            n += 1
        taken.add(name)
        return name

    for part in bundle.parts:
        for layer in part.geo_layers:
            name = claim(part.prefix + layer.name)
            out.append((name, layer.collection, OGR_GEOMETRY_TYPES[layer.question.input_type]))
        for results in part.results_layers:
            types = {(f.get('geometry') or {}).get('type') for f in results.collection.get('features') or []}
            types.discard(None)
            ogr_type = 'wkb' + types.pop() if len(types) == 1 else 'wkbUnknown'
            name = claim(part.prefix + 'layers_' + _sanitize_filename(results.layer_name) + '_results')
            out.append((name, results.collection, ogr_type))
    return out


def _write_vrt(bundle, workdir):
    """One OGR VRT naming every layer, so a single ogr2ogr call converts them
    all into one multi-layer target (GeoPackage, KML) or one directory (Shapefile)."""
    layers = []
    for name, collection, ogr_type in _unique_layer_names(bundle):
        collection = dict(collection, name=name)   # the GeoJSON driver reads the layer name from here
        src = os.path.join(workdir, name + '.geojson')
        with open(src, 'wb') as fh:
            fh.write(_geojson_bytes(collection))
        layers.append(
            '  <OGRVRTLayer name="{name}">\n'
            '    <SrcDataSource relativeToVRT="0">{src}</SrcDataSource>\n'
            '    <SrcLayer>{name}</SrcLayer>\n'
            '    <GeometryType>{gtype}</GeometryType>\n'
            '    <LayerSRS>WGS84</LayerSRS>\n'
            '  </OGRVRTLayer>\n'.format(name=xml_escape(name, {'"': '&quot;'}), src=xml_escape(src), gtype=ogr_type)
        )
    vrt = os.path.join(workdir, 'export.vrt')
    with open(vrt, 'w', encoding='utf-8') as fh:
        fh.write('<OGRVRTDataSource>\n' + ''.join(layers) + '</OGRVRTDataSource>\n')
    return vrt, len(layers)


def _ogr_formats():
    """Drivers this GDAL can write, cached per process."""
    if not hasattr(_ogr_formats, 'cache'):
        try:
            out = subprocess.run(['ogrinfo', '--formats'], capture_output=True, text=True, timeout=30).stdout
        except (OSError, subprocess.SubprocessError):
            out = ''
        _ogr_formats.cache = out
    return _ogr_formats.cache


def _run_ogr2ogr(args):
    try:
        result = subprocess.run(['ogr2ogr'] + args, capture_output=True, text=True, timeout=OGR_TIMEOUT)
    except subprocess.TimeoutExpired:
        raise ExportFailed(f'ogr2ogr timed out after {OGR_TIMEOUT}s')
    except OSError as exc:
        raise ExportFailed(f'ogr2ogr could not start: {exc}')
    if result.returncode != 0:
        raise ExportFailed(f'ogr2ogr exited {result.returncode}: {result.stderr.strip()}')
    if result.stderr.strip():
        # Shapefile field-name laundering lands here; the creator was told in the dialog.
        logger.info("ogr2ogr: %s", result.stderr.strip())
    return result


def write_ogr(bundle, fmt, workdir, base_name):
    """Convert every layer into ``fmt``. Returns the path of the single output
    file (gpkg, kml) or of the directory holding the shapefiles (shp)."""
    vrt, count = _write_vrt(bundle, workdir)
    if fmt == 'gpkg':
        target = os.path.join(workdir, base_name + '.gpkg')
        _run_ogr2ogr(['-f', 'GPKG', target, vrt])
        return target
    if fmt == 'kml':
        target = os.path.join(workdir, base_name + '.kml')
        driver = 'LIBKML' if 'LIBKML' in _ogr_formats() else 'KML'
        args = ['-f', driver, target, vrt]
        if driver == 'LIBKML':
            args += ['-dsco', 'NAME=' + base_name]
        _run_ogr2ogr(args)
        return target
    if fmt == 'shp':
        target = os.path.join(workdir, 'shp')
        os.makedirs(target)
        if count:
            _run_ogr2ogr(['-f', 'ESRI Shapefile', target, vrt, '-lco', 'ENCODING=UTF-8'])
        return target
    raise ValueError(fmt)


# ---------------------------------------------------------------------------
# Containers
# ---------------------------------------------------------------------------

def add_files_to_zip(bundle, zip):
    """Respondent uploads under files/<session_id>/, copied in chunks."""
    for part in bundle.parts:
        for answer in part.file_answers:
            arcname = part.prefix + _upload_archive_path(answer.question, answer)
            try:
                with answer.upload.file.open('rb') as src, zip.open(arcname, 'w') as dst:
                    shutil.copyfileobj(src, dst)
            except Exception:
                logger.warning(
                    "Export: upload %s missing from storage; row keeps the path, file absent from ZIP.",
                    answer.upload_id,
                )


def add_directory_to_zip(directory, zip):
    for root, _dirs, files in os.walk(directory):
        for name in sorted(files):
            full = os.path.join(root, name)
            zip.write(full, os.path.relpath(full, directory))


def build_export(bundle, fmt, include_files, workdir):
    """Write ``fmt`` into ``workdir``; return (path, download filename, content type).

    zip and shp are always an archive; xlsx, gpkg and kml are one file unless
    uploads are requested, in which case the file sits at the archive root next
    to files/; csv is an archive of the flat tables.
    """
    base = _sanitize_filename(bundle.survey.name)
    if fmt == 'xlsx':
        produced = write_xlsx(bundle, os.path.join(workdir, base + '.xlsx'))
    elif fmt in OGR_FORMATS:
        produced = write_ogr(bundle, fmt, workdir, base)
    else:
        produced = None

    if fmt in SINGLE_FILE_FORMATS and not include_files:
        return produced, os.path.basename(produced), CONTENT_TYPES.get(fmt, 'application/octet-stream')

    archive = os.path.join(workdir, base + '.export.zip')
    with ZipFile(archive, 'w', ZIP_DEFLATED) as zip:
        if fmt == 'csv':
            write_csv_zip(bundle, zip)
        elif fmt == 'shp':
            add_directory_to_zip(produced, zip)
        else:
            zip.write(produced, os.path.basename(produced))
        if include_files:
            add_files_to_zip(bundle, zip)
        for entry in zip.filelist:
            entry.create_system = 0   # Windows bug fix, as in the legacy archive
    return archive, base + '.zip', 'application/zip'


# ---------------------------------------------------------------------------
# Session filters
# ---------------------------------------------------------------------------

def excluded_sessions(version_surveys, include_all, completed_only):
    """Session ids kept out of the export: empty sessions always, trashed and
    not_approved unless include_all, plus every non-completed session when
    completed_only.
    "Completed" is the Responses overview's definition, through one helper."""
    from .analytics import completed_session_filter, empty_sessions
    headers = [header for header, _ in version_surveys]
    qs = SurveySession.objects.filter(survey__in=headers)
    # Sessions without a single answer are never exported, whatever
    # include_all and completed_only say (spec responses-export-formats):
    # they would only add blank rows.
    excluded = set(empty_sessions(qs).values_list('id', flat=True))
    if not include_all:
        from django.db.models import Q
        excluded |= set(
            qs.filter(Q(is_deleted=True) | Q(validation_status='not_approved'))
            .values_list('id', flat=True)
        )
    if completed_only:
        completed = set(completed_session_filter(qs, headers).values_list('id', flat=True))
        excluded |= set(qs.values_list('id', flat=True)) - completed
    return excluded
